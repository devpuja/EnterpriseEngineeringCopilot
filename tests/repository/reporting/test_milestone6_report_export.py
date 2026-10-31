from pathlib import Path
import base64
import html
import zipfile

import pytest

from app.repository.reporting.exporters.docx_exporter import DOCXExporter
from app.repository.reporting.exporters.html_exporter import HTMLExporter
from app.repository.reporting.exporters.markdown_exporter import MarkdownExporter
from app.repository.reporting.exporters.pdf_exporter import PDFExporter
from app.repository.reporting.models import ReportConfiguration
from app.repository.reporting.report_generator import ReportGenerator
from app.repository.reporting.validation import ReportValidationError, validate_report
from app.repository.review import ArchitectureReviewService
from app.repository.synthesis import ArchitectureSynthesizer
from app.repository.visualization.architecture_visualizer import ArchitectureVisualizer
from tests.repository.visualization.test_milestone3_visualization import analysis_fixture


def full_report(configuration: ReportConfiguration | None = None):
    analysis = analysis_fixture()
    visualization = ArchitectureVisualizer().visualize(analysis)
    synthesis = ArchitectureSynthesizer().synthesize(analysis)
    review = ArchitectureReviewService().review(synthesis, analysis)
    report = ReportGenerator().generate(
        analysis,
        visualization,
        configuration,
        synthesis=synthesis,
        review=review,
    )
    return analysis, synthesis, review, report


def test_canonical_report_integrates_all_milestones_and_eight_diagrams():
    analysis, synthesis, review, report = full_report()

    assert report.validated
    assert report.analysis_id == analysis.analysis_id
    assert report.synthesis_id == synthesis.synthesis_id
    assert report.review_id == review.review_id
    assert len(report.diagrams) == 8
    assert report.diagrams[-1].diagram_id == "unified_architecture"
    section_ids = [section.section_id for section in report.sections]
    assert section_ids == sorted(
        section_ids,
        key=lambda item: (
            "cover executive_summary repository_overview existing_architecture "
            "architecture_diagrams technology_stack components_dependencies api_architecture "
            "database_architecture external_integrations architecture_findings architecture_synthesis "
            "architecture_components architecture_relationships architecture_data_flows external_systems "
            "architecture_review review_findings architecture_strengths recommendations limitations "
            "evidence_references appendix"
        ).split().index(item),
    )
    assert any("Overall outcome:" in item.content for item in report.sections)
    assert report.synthesis.components
    assert report.review.findings or report.review.strengths


def test_report_generation_rejects_mismatched_milestone_inputs():
    analysis = analysis_fixture()
    visualization = ArchitectureVisualizer().visualize(analysis)
    synthesis = ArchitectureSynthesizer().synthesize(analysis)
    synthesis.original_analysis_id = analysis_fixture().analysis_id

    with pytest.raises(ValueError, match="synthesis.original_analysis_id"):
        ReportGenerator().generate(analysis, visualization, synthesis=synthesis)

    with pytest.raises(ReportValidationError, match="unknown section IDs"):
        ReportGenerator().generate(
            analysis,
            visualization,
            ReportConfiguration(include_sections=["not_a_section"]),
        )
    with pytest.raises(ReportValidationError, match="cannot be both included and excluded"):
        ReportGenerator().generate(
            analysis,
            visualization,
            ReportConfiguration(
                include_sections=["api_architecture"],
                exclude_sections=["api_architecture"],
            ),
        )
    with pytest.raises(ValueError, match="analysis result is required"):
        ReportGenerator().generate(None, visualization)


def test_invalid_evidence_reference_fails_report_validation():
    _, _, _, report = full_report()
    report.sections[0].evidence_references.append("UNVERIFIED")

    with pytest.raises(ReportValidationError, match="unknown evidence IDs UNVERIFIED"):
        validate_report(report, require_marked=True)


def test_report_configuration_is_applied_before_each_export(tmp_path):
    configuration = ReportConfiguration(
        selected_diagram_types=["api", "unified_architecture"],
        include_evidence=False,
        include_confidence=False,
        include_limitations=False,
        include_mermaid_source=False,
        include_rendered_diagrams=False,
    )
    _, _, _, report = full_report(configuration)

    assert any(item.diagram_id == "unified_architecture" for item in report.diagrams)
    assert any(item.diagram_type.value == "api" for item in report.diagrams)
    assert all(
        item.diagram_id == "unified_architecture" or item.diagram_type.value == "api"
        for item in report.diagrams
    )
    assert not report.evidence
    assert all(not item.mermaid_source for item in report.diagrams)
    assert all(not item.rendered_png and not item.rendered_svg for item in report.diagrams)
    assert "limitations" not in {item.section_id for item in report.sections}
    assert all(not item.evidence_references for item in report.sections)

    results = [
        MarkdownExporter().export(report, tmp_path / "md"),
        HTMLExporter().export(report, tmp_path / "html"),
        PDFExporter().export(report, tmp_path / "pdf"),
        DOCXExporter().export(report, tmp_path / "docx"),
    ]
    assert all(item.export_status == "completed" for item in results)
    assert len({item.report_id for item in results}) == 1
    markdown = (tmp_path / "md" / str(report.report_id) / results[0].output_file_name).read_text(encoding="utf-8")
    html_report = (tmp_path / "html" / str(report.report_id) / results[1].output_file_name).read_text(encoding="utf-8")
    assert "```mermaid" not in markdown
    assert "Evidence:" not in markdown
    assert "confidence" not in markdown.lower()
    assert "flowchart LR" not in html_report
    assert "Evidence:" not in html_report

    from pypdf import PdfReader
    from docx import Document

    pdf_path = tmp_path / "pdf" / str(report.report_id) / results[2].output_file_name
    docx_path = tmp_path / "docx" / str(report.report_id) / results[3].output_file_name
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages)
    docx_text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    for content in (pdf_text, docx_text):
        assert "Evidence:" not in content
        assert "confidence" not in content.lower()
        assert "flowchart LR" not in content


def test_exporters_require_validated_model_and_preserve_shared_content(tmp_path):
    _, _, review, report = full_report()
    invalid_report = report.model_copy(update={"validated": False})
    failed = MarkdownExporter().export(invalid_report, tmp_path / "invalid")
    assert failed.export_status == "failed"
    assert "validated report model" in failed.warnings[0]

    results = [
        MarkdownExporter().export(report, tmp_path / "md"),
        HTMLExporter().export(report, tmp_path / "html"),
        PDFExporter().export(report, tmp_path / "pdf"),
        DOCXExporter().export(report, tmp_path / "docx"),
    ]
    assert all(item.export_status == "completed" for item in results)

    md_path = tmp_path / "md" / str(report.report_id) / results[0].output_file_name
    html_path = tmp_path / "html" / str(report.report_id) / results[1].output_file_name
    pdf_path = tmp_path / "pdf" / str(report.report_id) / results[2].output_file_name
    docx_path = tmp_path / "docx" / str(report.report_id) / results[3].output_file_name
    assert all(Path(path).is_file() for path in (md_path, html_path, pdf_path, docx_path))
    assert pdf_path.read_bytes().startswith(b"%PDF-")
    assert zipfile.is_zipfile(docx_path)

    markdown = md_path.read_text(encoding="utf-8")
    html_text = html.unescape(html_path.read_text(encoding="utf-8"))
    from pypdf import PdfReader

    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages)
    from docx import Document

    docx_text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    marker = f"Overall outcome: {review.overall_review_outcome.value}"
    assert marker in markdown
    assert marker in html_text
    assert marker in pdf_text
    assert marker in docx_text
    assert review.overall_review_outcome.value in markdown


def test_markdown_export_embeds_rendered_diagram_with_relative_artifact(tmp_path):
    _, _, _, report = full_report()
    report.diagrams[0].rendered_png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    result = MarkdownExporter().export(report, tmp_path / "md")

    assert result.export_status == "completed"
    report_directory = tmp_path / "md" / str(report.report_id)
    content = (report_directory / result.output_file_name).read_text(encoding="utf-8")
    asset_name = f"diagram-{report.diagrams[0].diagram_id}.png"
    assert f"]({asset_name})" in content
    assert (report_directory / asset_name).is_file()
    assert result.output_location == result.managed_artifact_reference


def test_repository_secrets_and_absolute_paths_are_redacted():
    analysis = analysis_fixture()
    analysis.source_type = "local"
    analysis.source_reference = r"C:\private\repository"
    analysis.summary = "token=do-not-export from C:\\private\\repository"
    visualization = ArchitectureVisualizer().visualize(analysis)
    visualization.diagrams[0].mermaid_source += "\n%% password=diagram-secret"
    visualization.diagrams[0].rendered_png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    report = ReportGenerator().generate(
        analysis,
        visualization,
        ReportConfiguration(organization_name="password=do-not-export"),
    )
    serialized = report.model_dump_json()

    assert "do-not-export" not in serialized
    assert "diagram-secret" not in serialized
    assert r"C:\private\repository" not in serialized
    assert "Managed local repository" in report.repository_reference
    assert report.diagrams[0].rendered_png is None


def test_exporter_rejects_artifacts_inside_local_source_repository(tmp_path):
    source_root = tmp_path / "source"
    source_root.mkdir()
    analysis = analysis_fixture()
    analysis.source_type = "local"
    analysis.source_reference = str(source_root)
    report = ReportGenerator().generate(
        analysis,
        ArchitectureVisualizer().visualize(analysis),
    )
    output_directory = source_root / "reports"

    result = MarkdownExporter().export(report, output_directory)

    assert result.export_status == "failed"
    assert "must not be inside the source repository" in result.warnings[0]
    assert not output_directory.exists()
    assert str(source_root) not in result.model_dump_json()
