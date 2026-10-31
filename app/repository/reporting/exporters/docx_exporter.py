from pathlib import Path
from io import BytesIO

from app.repository.reporting.artifact_manager import ArtifactManager
from app.repository.reporting.exporters.base import (
    failed_export,
    get_artifact_manager,
    redact_sensitive_text,
    validate_export_report,
    valid_png,
)
from app.repository.reporting.models import ArchitectureReport, ReportExportResult, ReportFormat


class DOCXExporter:
    def __init__(self, artifact_manager: ArtifactManager | None = None):
        self.artifact_manager = artifact_manager

    def export(self, report: ArchitectureReport, output_directory: Path) -> ReportExportResult:
        try:
            validate_export_report(report)
        except ValueError as exc:
            return failed_export(report, ReportFormat.docx, "docx", f"Invalid report: {exc}")

        warnings: list[str] = []
        try:
            manager = get_artifact_manager(report, output_directory, self.artifact_manager)
        except (OSError, ValueError) as exc:
            return failed_export(
                report,
                ReportFormat.docx,
                "docx",
                f"Unsafe artifact output path: {type(exc).__name__}: {exc}",
            )
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Inches, Pt
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError:
            return ReportExportResult(
                report_id=report.report_id,
                output_format=ReportFormat.docx,
                output_file_name=manager.safe_filename(report.report_id, "docx"),
                export_status="failed",
                warnings=["DOCX export unavailable: python-docx dependency is not installed."],
            )

        try:
            document = Document()
            section = document.sections[0]
            section.header.paragraphs[0].text = redact_sensitive_text(report.repository_name)
            footer = section.footer.paragraphs[0]
            footer.add_run(redact_sensitive_text(report.title) + " | Page ")
            _append_page_number(footer, OxmlElement, qn)

            title = document.add_paragraph()
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.add_run(redact_sensitive_text(report.title)).bold = True
            title.runs[0].font.size = Pt(24)
            document.add_paragraph("Repository: " + redact_sensitive_text(report.repository_name))
            document.add_paragraph("Source type: " + redact_sensitive_text(report.repository_source_type))
            document.add_paragraph("Source reference: " + redact_sensitive_text(report.repository_reference))
            document.add_paragraph(report.report_metadata.generated_at)
            document.add_paragraph("Report version: " + report.report_metadata.report_version)
            document.add_page_break()

            document.add_heading("Table of Contents", level=1)
            for report_section in report.sections:
                document.add_paragraph(
                    redact_sensitive_text(report_section.title),
                    style="List Bullet",
                )
            document.add_page_break()

            for report_section in report.sections:
                document.add_heading(redact_sensitive_text(report_section.title), level=1)
                for line in redact_sensitive_text(report_section.content).splitlines():
                    if line.startswith("- "):
                        document.add_paragraph(line[2:], style="List Bullet")
                    else:
                        document.add_paragraph(line)
                if report.configuration.include_evidence and report_section.evidence_references:
                    document.add_paragraph(
                        "Evidence: " + ", ".join(report_section.evidence_references)
                    )

            for diagram in report.diagrams:
                document.add_heading(redact_sensitive_text(diagram.title), level=1)
                if report.configuration.include_rendered_diagrams and diagram.rendered_png:
                    try:
                        if not valid_png(diagram.rendered_png):
                            raise ValueError("Invalid PNG content.")
                        document.add_picture(BytesIO(diagram.rendered_png), width=Inches(6.0))
                    except Exception as exc:
                        warnings.append(
                            f"Rendered image for diagram {diagram.diagram_id} could not be embedded: "
                            f"{type(exc).__name__}."
                        )
                elif report.configuration.include_rendered_diagrams and diagram.rendered_svg:
                    warnings.append(
                        f"SVG rendering for diagram {diagram.diagram_id} is not supported by the DOCX exporter."
                    )
                elif report.configuration.include_rendered_diagrams:
                    warnings.append(f"No rendered image is available for diagram {diagram.diagram_id}.")
                if report.configuration.include_mermaid_source:
                    document.add_paragraph("Mermaid source fallback:")
                    document.add_paragraph(redact_sensitive_text(diagram.mermaid_source), style="No Spacing")
                if report.configuration.include_evidence and diagram.evidence_references:
                    document.add_paragraph("Evidence: " + ", ".join(diagram.evidence_references))

            document.core_properties.title = redact_sensitive_text(report.title)
            document.core_properties.subject = "Architecture analysis report"
            buffer = BytesIO()
            document.save(buffer)
            content = buffer.getvalue()
        except Exception as exc:
            return ReportExportResult(
                report_id=report.report_id,
                output_format=ReportFormat.docx,
                output_file_name=manager.safe_filename(report.report_id, "docx"),
                export_status="failed",
                warnings=[f"DOCX export failed: {type(exc).__name__}: {redact_sensitive_text(str(exc))}"],
            )

        try:
            _, relative, size = manager.write_bytes(report.report_id, "docx", content)
        except (OSError, ValueError) as exc:
            return ReportExportResult(
                report_id=report.report_id,
                output_format=ReportFormat.docx,
                output_file_name=manager.safe_filename(report.report_id, "docx"),
                export_status="failed",
                warnings=[f"DOCX artifact write failed: {type(exc).__name__}: {redact_sensitive_text(str(exc))}"],
            )
        return ReportExportResult(
            report_id=report.report_id,
            output_format=ReportFormat.docx,
            output_file_name=Path(relative).name,
            managed_artifact_reference=relative,
            output_location=relative,
            file_size=size,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            export_status="completed",
            warnings=[redact_sensitive_text(item) for item in warnings],
        )


def _append_page_number(paragraph, element_factory, qualified_name) -> None:
    field = element_factory("w:fldSimple")
    field.set(qualified_name("w:instr"), "PAGE")
    paragraph._p.append(field)
